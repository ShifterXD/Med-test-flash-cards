from pathlib import Path
import csv
import json
import re
from docx import Document

ROOT = Path(__file__).resolve().parent
ATT = Path('/root/.hermes/kanban/attachments/t_272aa434')
DOCX = ATT / 'FSP Stuttgart.docx'

if not DOCX.exists():
    raise SystemExit(f'Missing source DOCX: {DOCX}')

extracted = ROOT / 'extracted'
extracted.mkdir(exist_ok=True)

doc = Document(str(DOCX))
lines = []
for p in doc.paragraphs:
    text = p.text.strip()
    if text:
        lines.append(text)
for table_index, table in enumerate(doc.tables, start=1):
    lines.append(f'\n[TABLE {table_index}]')
    for row in table.rows:
        cells = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
        if any(cells):
            lines.append(' || '.join(cells))
(extracted / 'FSP Stuttgart.txt').write_text('\n'.join(lines), encoding='utf-8')

cards = []
current_table = None
for line in lines:
    table_match = re.match(r'\[TABLE (\d+)\]', line.strip())
    if table_match:
        current_table = int(table_match.group(1))
        continue
    if '||' not in line:
        continue
    parts = [part.strip().replace(' | ', ' ').replace('| ', ' ').replace(' |', ' ') for part in line.split('||')]
    if len(parts) < 3:
        continue
    frequency_text, term, meaning = parts[0], parts[1], ' || '.join(parts[2:])
    if not frequency_text.isdigit() or not term or not meaning:
        continue
    frequency = int(frequency_text)
    difficulty = 'high-priority' if frequency >= 5 else 'medium' if frequency >= 3 else 'basic'
    term = re.sub(r'\s+', ' ', term).strip()
    meaning = re.sub(r'\s+', ' ', meaning).strip()
    cards.append({
        'id': f'fsp-stuttgart-{len(cards) + 1:03d}',
        'source': 'FSP Stuttgart.docx',
        'source_table': current_table,
        'topic': 'Medizinische Fachbegriffe FSP Stuttgart',
        'category': 'Medizinische Fachbegriffe FSP Stuttgart',
        'difficulty': difficulty,
        'frequency': frequency,
        'direction': 'Fachbegriff → Umgangssprache',
        'term': term,
        'question': f'Was bedeutet der medizinische Fachbegriff „{term}“ in Umgangssprache?',
        'answer': meaning,
    })

if not cards:
    raise SystemExit('No cards were extracted from DOCX tables')

schema = {
    'format_version': '1.0',
    'description': 'Flashcards for FSP Stuttgart medical terminology. The app requires id/topic/question/answer; additional fields are preserved for filtering and metadata.',
    'fields': {
        'id': 'Stable card identifier',
        'topic': 'Topic/category for filtering',
        'question': 'Prompt shown on the front of the card',
        'answer': 'Expected answer shown after reveal',
        'difficulty': 'basic | medium | high-priority, derived from Häufigkeit',
        'frequency': 'Häufigkeit value from source table',
        'term': 'Medical term',
        'source': 'Original source filename'
    }
}

(ROOT / 'cards.json').write_text(json.dumps({'cards': cards}, ensure_ascii=False, indent=2), encoding='utf-8')
(ROOT / 'cards.schema.json').write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding='utf-8')
with (ROOT / 'cards.csv').open('w', encoding='utf-8', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=list(cards[0].keys()))
    writer.writeheader()
    writer.writerows(cards)

print(f'Generated {len(cards)} cards from {DOCX}')
print(f'First: {cards[0]["question"]} => {cards[0]["answer"]}')
print(f'Last: {cards[-1]["question"]} => {cards[-1]["answer"]}')
